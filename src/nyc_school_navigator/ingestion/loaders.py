from pathlib import Path

from openpyxl import load_workbook


TARGET_SCHOOLS = {
    "20K187": "The Christa McAuliffe School / I.S. 187",
    "21K228": "I.S. 228 David A. Boody",
    "21K239": "Mark Twain I.S. 239 for the Gifted & Talented",
}


SUMMARY_FIELDS = {
    "DBN": "dbn",
    "School Name": "school_name",
    "Enrollment": "enrollment",
    "Instruction and Performance - Rating": "instruction_performance_rating",
    "Safety and School Climate - Rating": "safety_school_climate_rating",
    "Relationships with Families - Rating": "relationships_with_families_rating",
    "Instruction/Learning Environment - School Percent Positive": (
        "learning_environment_positive"
    ),
    "Safety - School Percent Positive": "safety_positive",
    "Student Support - School Percent Positive": "student_support_positive",
    "Advising and Planning - School Percent Positive": "advising_planning_positive",
    "Family-School Trust - School Percent Positive": "family_school_trust_positive",
    "Communication - School Percent Positive": "communication_positive",
    "Family Involvement - School Percent Positive": "family_involvement_positive",
    "Average Student Attendance": "average_student_attendance",
}


ADDITIONAL_FIELDS = {
    "DBN": "dbn",
    "Metric Value - % of 8th Graders earning High School Credit in Math": (
        "high_school_math_credit"
    ),
    "Metric Value - % of 8th Graders earning High School Credit in Science": (
        "high_school_science_credit"
    ),
}

def _extract_rows(
    worksheet,
    fields: dict[str, str],
) -> dict[str, dict]:
    """Extract selected fields for the target schools."""
    rows = worksheet.iter_rows(values_only=True)

    # Row 4 contains headers
    for _ in range(3):
        next(rows)

    header_row = next(rows)

    headers = {
        str(value).strip(): index
        for index, value in enumerate(header_row)
        if value is not None
    }

    missing_fields = [field for field in fields if field not in headers]

    if missing_fields:
        raise ValueError(
            f"Missing expected columns in sheet '{worksheet.title}': "
            f"{missing_fields}"
        )

    schools = {}

    # Row 5 is not data, so skip it
    next(rows)

    for row in rows:
        dbn = row[headers["DBN"]]

        if dbn not in TARGET_SCHOOLS:
            continue

        record = {
            normalized_field: row[headers[source_field]]
            for source_field, normalized_field in fields.items()
        }

        schools[dbn] = record

    return schools


def load_school_quality(workbook_path: str | Path) -> list[dict]:
    """Load MVP school-quality data for the three selected schools."""
    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=True,
    )

    summary = _extract_rows(
        workbook["Summary"],
        SUMMARY_FIELDS,
    )

    additional = _extract_rows(
        workbook["Additional Info"],
        ADDITIONAL_FIELDS,
    )

    missing_schools = TARGET_SCHOOLS.keys() - summary.keys()

    if missing_schools:
        raise ValueError(
            f"Target schools missing from Summary sheet: "
            f"{sorted(missing_schools)}"
        )

    results = []

    for dbn in TARGET_SCHOOLS:
        record = summary[dbn]

        if dbn in additional:
            record.update(
                {
                    key: value
                    for key, value in additional[dbn].items()
                    if key != "dbn"
                }
            )

        record["source_type"] = "school_quality"
        record["source_year"] = "2024-25"

        results.append(record)

    return results